import json
import os
import ssl
import warnings
import Notification
from datetime import date
from pathlib import Path
from urllib.parse import quote, urlencode
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


class DepartmentMappingError(ValueError):
	pass


class EmployeeDataError(ValueError):
	pass


def load_env_file():
	env_path = Path(__file__).with_name(".env")
	for line in env_path.read_text(encoding="utf-8").splitlines():
		line = line.strip()
		if not line or line.startswith("#") or "=" not in line:
			continue
		name, value = line.split("=", 1)
		os.environ[name.strip()] = value.strip().strip("\"'")


def create_ssl_context():
	ca_bundle = os.environ.get("MDM_CA_BUNDLE")
	context = ssl.create_default_context(cafile=ca_bundle)
	if os.name == "nt" and not ca_bundle:
		certificates = []
		for store_name in ("ROOT", "CA"):
			for certificate, encoding, trust in ssl.enum_certificates(store_name):
				trusted_for_server = trust is True or "1.3.6.1.5.5.7.3.1" in trust
				if encoding == "x509_asn" and trusted_for_server:
					certificates.append(ssl.DER_cert_to_PEM_cert(certificate))
		if certificates:
			context.load_verify_locations(cadata="".join(certificates))
	return context


def load_processed_user_ids(path):
	if not path.exists():
		path.write_text("{}\n", encoding="utf-8")
	contents = path.read_text(encoding="utf-8")
	if not contents.strip():
		processed_user_ids = {}
		save_processed_user_ids(path, processed_user_ids)
	else:
		processed_user_ids = json.loads(contents)
	if isinstance(processed_user_ids, list):
		processed_user_ids = {
			str(user_id): {
				"EmployeeUpdateSuccess": True,
				"EmployeeNFCUpdateSuccess": True,
			}
			for user_id in processed_user_ids
		}
		save_processed_user_ids(path, processed_user_ids)
		return processed_user_ids
	if not isinstance(processed_user_ids, dict):
		raise ValueError("ProcessedUserID.json must contain a JSON object")
	statuses = {}
	for user_id, user_status in processed_user_ids.items():
		if not isinstance(user_status, dict):
			raise ValueError(f"Processed status for {user_id} must be a JSON object")
		statuses[str(user_id)] = {}
		for status_name in ("EmployeeUpdateSuccess", "EmployeeNFCUpdateSuccess"):
			status = user_status.get(status_name, False)
			if not isinstance(status, bool):
				raise ValueError(f"{status_name} for {user_id} must be a boolean")
			statuses[str(user_id)][status_name] = status
	return statuses


def save_processed_user_ids(path, processed_user_ids):
	path.write_text(
		json.dumps(processed_user_ids, indent=2) + "\n",
		encoding="utf-8",
	)


def notify_api_error(message):
	try:
		Notification.send_error_notification(
			subject="MDM 2.0 to RMS API Error",
			body=str(message),
		)
	except Exception as notification_error:
		print(f"Failed to send error notification: {notification_error}")


def report_employee_error(user_id, action, error):
	message = f"{action} for employee {user_id} failed: {error}"
	if isinstance(error, HTTPError):
		response_body = error.read().decode("utf-8", errors="replace")
		if response_body:
			message += f"\nResponse: {response_body}"
	print(message)
	notify_api_error(message)


def get_rms_token(settings, ssl_context):
	token_request = Request(
		settings["RMS_TOKEN_URL"],
		data=json.dumps(
			{
				"UserID": settings["RMS_USER"],
				"Password": settings["RMS_PASSWORD"],
			}
		).encode("utf-8"),
		headers={"Content-Type": "application/json"},
		method="POST",
	)
	with urlopen(token_request, context=ssl_context) as response:
		token_payload = json.load(response)
	issued_token = token_payload.get("IssuedToken")
	if not isinstance(issued_token, str) or not issued_token:
		raise ValueError("RMS token response did not contain IssuedToken")
	return issued_token


def rms_employee_exists(user_id, employee_url, issued_token, ssl_context):
	retrieve_request = Request(
		f"{employee_url.rstrip('/')}/{quote(str(user_id), safe='')}",
		headers={"Token": issued_token, "Accept": "application/json"},
		method="GET",
	)
	try:
		with urlopen(retrieve_request, context=ssl_context):
			return True
	except HTTPError as error:
		if error.code in {400, 404}:
			return False
		raise


def load_department_mappings(mapping_path):
	with warnings.catch_warnings():
		warnings.filterwarnings(
			"ignore",
			message="Data Validation extension is not supported and will be removed",
			category=UserWarning,
			module=r"openpyxl\.worksheet\._reader",
		)
		workbook = load_workbook(mapping_path, data_only=True)
	department_mappings = {}
	mapping_table_found = False
	try:
		for worksheet in workbook.worksheets:
			for table in worksheet.tables.values():
				min_column, header_row, max_column, max_row = range_boundaries(table.ref)
				headers = {
					str(worksheet.cell(header_row, column).value).strip(): column
					for column in range(min_column, max_column + 1)
					if worksheet.cell(header_row, column).value is not None
				}
				if not {"MDM Department", "Department Code"}.issubset(headers):
					continue
				if mapping_table_found:
					raise ValueError("Multiple tables contain MDM Department and Department Code")
				mapping_table_found = True
				for row in range(header_row + 1, max_row + 1):
					department_name = worksheet.cell(row, headers["MDM Department"]).value
					department_code = worksheet.cell(row, headers["Department Code"]).value
					if department_name is None or not str(department_name).strip():
						continue
					if department_code is None or not str(department_code).strip():
						raise ValueError(
							f"Missing Department Code for MDM Department {department_name!r}"
						)
					department_name = str(department_name).strip()
					department_code = str(department_code).strip()
					previous_code = department_mappings.get(department_name)
					if previous_code is not None and previous_code != department_code:
						raise ValueError(
							f"Conflicting Department Codes for MDM Department {department_name!r}"
						)
					department_mappings[department_name] = department_code
	finally:
		workbook.close()
	if not mapping_table_found:
		raise ValueError(
			"Mapping.xlsx must contain a table with MDM Department and Department Code columns"
		)
	return department_mappings


def build_rms_employee(employee, department_mappings):
	accounts = employee.get("Accounts") or []
	sam_account_name = next(
		(
			account.get("SamAccountName")
			for account in accounts
			if account.get("SamAccountName") and account.get("IsPrimary") and account.get("IsActive")
		),
		None,
	)
	if sam_account_name is None:
		sam_account_name = next(
			(account.get("SamAccountName") for account in accounts if account.get("SamAccountName")),
			None,
		)
	phones = employee.get("Phones") or []
	emails = employee.get("Emails") or []
	phone_number = next(
		(phone.get("PhoneNumber") for phone in phones if phone.get("PhoneNumber")),
		None,
	)
	active_primary_email = next(
		(
			email.get("Email")
			for email in emails
			if email.get("Email") and email.get("IsPrimary") and email.get("IsActive")
		),
		None,
	)
	department_name = (employee.get("PrimaryDepartment") or {}).get("Name")
	if not isinstance(department_name, str) or not department_name.strip():
		raise EmployeeDataError(
			f"Employee {employee.get('EmployeeNo')} is missing PrimaryDepartment.Name"
		)
	department_name = department_name.strip()
	try:
		department_code = department_mappings[department_name]
	except KeyError as error:
		raise DepartmentMappingError(
			f"No Department Code mapping found for MDM Department {department_name!r}"
		) from error
	ad_domain_prefix = "NFTRINITY" if department_code == "NTR" else "NANFUNG"
	email_prefix = "" if "prod" in os.environ.get("ENVIRONMENT", "").casefold() else "dev."
	new_employee = {
		"Revision": 1,
		"UserID": employee["EmployeeNo"],
		"UserNameChi": "",
		"UserNameEng": employee.get("PreferedName") or employee.get("PreferredName") or "",
		"ADUserID": f"{ad_domain_prefix}\\{sam_account_name or ''}",
		"Department": department_code,
		"Tel": phone_number or "",
		"Email": f"{email_prefix}{active_primary_email}" if active_primary_email else "",
		"UserGroup": "NonUser",
		"IsActive": 1,
	}
	return new_employee


def create_rms_employee(rms_employee, employee_url, issued_token, ssl_context):
	create_request = Request(
		employee_url.rstrip("/"),
		data=json.dumps([rms_employee]).encode("utf-8"),
		headers={
			"Token": issued_token,
			"Content-Type": "application/json",
			"Accept": "application/json",
		},
		method="POST",
	)
	with urlopen(create_request, context=ssl_context) as response:
		return response.status


def create_rms_employee_nfc(employee, employee_nfc_url, issued_token, ssl_context):
	staff_cards = employee.get("StaffCards") or []
	nfc_id = next(
		(
			card.get("StaffCardId")
			for card in staff_cards
			if card.get("StaffCardId") is not None
			and str(card.get("StaffCardId")).strip()
		),
		None,
	)
	if nfc_id is None:
		return None
	nfc_record = {
		"Revision": 1,
		"UserID": employee["EmployeeNo"],
		"NFCs": [{"NFCID": nfc_id}],
	}
	nfc_request = Request(
		employee_nfc_url.rstrip("/"),
		data=json.dumps([nfc_record]).encode("utf-8"),
		headers={
			"Token": issued_token,
			"Content-Type": "application/json",
			"Accept": "application/json",
		},
		method="POST",
	)
	with urlopen(nfc_request, context=ssl_context) as response:
		return response.status


def _process():
	load_env_file()
	token_url = os.environ["MDM_TOKEN_URL"]
	api_key = os.environ["MDM_KEY"]
	employee_url = os.environ["MDM_EMPLOYEE_URL"]
	ssl_context = create_ssl_context()

	token_request = Request(token_url, headers={"x-api-key": api_key})
	with urlopen(token_request, context=ssl_context) as response:
		access_token = json.load(response)["AccessToken"]

	query_size = int(os.environ["MDM_QUERY_SIZE"])
	if query_size < 1:
		raise ValueError("MDM_QUERY_SIZE must be a positive integer")
	separator = "&" if "?" in employee_url else "?"
	employee_list = None
	all_employees = []
	page = 1
	while True:
		query = urlencode({"Page": page, "Size": query_size})
		employee_request = Request(
			f"{employee_url}{separator}{query}",
			headers={"Authorization": f"Bearer {access_token}"},
		)
		print(f"Requesting MDM employee page {page}")
		with urlopen(employee_request, context=ssl_context) as response:
			page_data = json.load(response)
		if employee_list is None:
			employee_list = page_data
		page_results = page_data.get("Results", [])
		all_employees.extend(page_results)
		total_pages = page_data.get("TotalPages")
		if not page_results or (total_pages is not None and page >= int(total_pages)):
			break
		page += 1
	employee_list["Results"] = all_employees
	cutoff_date = date.fromisoformat(os.environ["MDM_ORIGINAL_HIRE_DATE"])
	filtered_employees = []
	for employee in employee_list["Results"]:
		user_id = employee.get("EmployeeNo", "unknown") if isinstance(employee, dict) else "unknown"
		if not isinstance(employee, dict):
			report_employee_error(
				user_id,
				"Employee data validation",
				EmployeeDataError("MDM employee result is not an object"),
			)
			continue
		hire_date = employee.get("OriginalHireDate")
		if not hire_date:
			continue
		try:
			original_hire_date = date.fromisoformat(hire_date[:10])
		except Exception as error:
			report_employee_error(user_id, "Employee hire-date validation", error)
			continue
		if original_hire_date > cutoff_date:
			filtered_employees.append(employee)
	employee_list["Results"] = filtered_employees
	employee_list["Count"] = len(employee_list["Results"])
	department_mappings = load_department_mappings(
		Path(__file__).with_name("Mapping.xlsx")
	)
	rms_update_enabled = os.environ.get("RMS_UPDATE_ENABLED", "false").strip().lower()
	if rms_update_enabled not in {"true", "false"}:
		raise ValueError("RMS_UPDATE_ENABLED must be true or false")
	processed_ids_path = Path(__file__).with_name("ProcessedUserID.json")
	processed_user_ids = load_processed_user_ids(processed_ids_path)
	rms_employees = []
	rms_lookup_token = None
	rms_employee_url = os.environ.get("RMS_EMPLOYEE_URL")
	if rms_update_enabled == "true" and not rms_employee_url:
		raise ValueError("RMS_EMPLOYEE_URL is required when RMS_UPDATE_ENABLED is true")
	if rms_update_enabled == "true" and employee_list["Results"]:
		try:
			rms_lookup_token = get_rms_token(os.environ, ssl_context)
		except Exception as error:
			message = f"RMS employee lookup token request failed: {error}"
			print(message)
			notify_api_error(message)
	for employee in employee_list["Results"]:
		user_id = employee.get("EmployeeNo", "unknown") if isinstance(employee, dict) else "unknown"
		try:
			rms_employee = build_rms_employee(employee, department_mappings)
		except Exception as error:
			if isinstance(error, DepartmentMappingError):
				action = "Department mapping"
			elif isinstance(error, EmployeeDataError):
				action = "Employee data validation"
			else:
				action = "Employee processing"
			report_employee_error(user_id, action, error)
			continue
		if rms_update_enabled == "true":
			if rms_lookup_token is None:
				continue
			try:
				exists_in_rms = rms_employee_exists(
					employee["EmployeeNo"],
					rms_employee_url,
					rms_lookup_token,
					ssl_context,
				)
			except Exception as error:
				report_employee_error(user_id, "RMS employee lookup", error)
				continue
			user_id = str(employee["EmployeeNo"])
			user_status = processed_user_ids.setdefault(
				user_id,
				{
					"EmployeeUpdateSuccess": False,
					"EmployeeNFCUpdateSuccess": False,
				},
			)
			status_changed = user_status["EmployeeUpdateSuccess"] != exists_in_rms
			user_status["EmployeeUpdateSuccess"] = exists_in_rms
			if exists_in_rms and not user_status["EmployeeNFCUpdateSuccess"]:
				user_status["EmployeeNFCUpdateSuccess"] = True
				status_changed = True
			if status_changed:
				save_processed_user_ids(processed_ids_path, processed_user_ids)
		rms_employees.append((employee, rms_employee))

	pending_employees = [
		(employee, rms_employee)
		for employee, rms_employee in rms_employees
		if employee.get("EmployeeNo") is not None
		and not all(processed_user_ids.get(str(employee["EmployeeNo"]), {}).get(status_name, False)
				for status_name in ("EmployeeUpdateSuccess", "EmployeeNFCUpdateSuccess"))
	]
	for employee, rms_employee in pending_employees:
		user_id = employee.get("EmployeeNo", "unknown")
		try:
			employee_summary = dict(rms_employee)
			employee_summary["NFCIDs"] = [
				card.get("StaffCardId")
				for card in employee.get("StaffCards") or []
				if card.get("StaffCardId")
			]
			print(json.dumps(employee_summary, ensure_ascii=False, indent=2))
		except Exception as error:
			report_employee_error(user_id, "Employee summary", error)
	if rms_update_enabled == "true" and pending_employees:
		try:
			rms_token = get_rms_token(os.environ, ssl_context)
		except Exception as error:
			message = f"RMS update token request failed; employee updates were skipped: {error}"
			print(message)
			notify_api_error(message)
			pending_employees = []
	if rms_update_enabled == "true" and pending_employees:
		rms_employee_url = os.environ["RMS_EMPLOYEE_URL"]
		rms_employee_nfc_url = os.environ["RMS_EMPLOYEE_NFC_URL"]
		for employee, rms_employee in pending_employees:
			user_id = str(employee["EmployeeNo"])
			user_status = processed_user_ids.setdefault(
				user_id,
				{
					"EmployeeUpdateSuccess": False,
					"EmployeeNFCUpdateSuccess": False,
				},
			)
			if not user_status["EmployeeUpdateSuccess"]:
				try:
					status = create_rms_employee(
						rms_employee, rms_employee_url, rms_token, ssl_context
					)
				except Exception as error:
					print(f"Create employee {user_id}: FAILED ({error})")
					report_employee_error(user_id, "RMS employee creation", error)
					continue
				print(f"Create employee {user_id}: SUCCESS (HTTP {status})")
				user_status["EmployeeUpdateSuccess"] = True
				save_processed_user_ids(processed_ids_path, processed_user_ids)
			if not user_status["EmployeeNFCUpdateSuccess"]:
				try:
					nfc_status = create_rms_employee_nfc(
						employee, rms_employee_nfc_url, rms_token, ssl_context
					)
				except Exception as error:
					print(f"Create employee NFC {user_id}: FAILED ({error})")
					report_employee_error(user_id, "RMS EmployeeNFC creation", error)
					continue
				if nfc_status is None:
					print(f"Create employee NFC {user_id}: SKIPPED (no StaffCardId in MDM)")
				else:
					print(f"Create employee NFC {user_id}: SUCCESS (HTTP {nfc_status})")
				user_status["EmployeeNFCUpdateSuccess"] = True
				save_processed_user_ids(processed_ids_path, processed_user_ids)
	elif rms_update_enabled == "false":
		print("RMS updates disabled (RMS_UPDATE_ENABLED=false)")

	output_path = Path(__file__).with_name("employeeList.json")
	output_path.write_text(
		json.dumps(employee_list, indent=2, ensure_ascii=False) + "\n",
		encoding="utf-8",
	)


def main():
	try:
		_process()
	except Exception as error:
		notify_api_error(f"API processing failed: {error}")
		raise


if __name__ == "__main__":
	main()
